import re
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill
import os
import sys

def parse_sql_file(sql_file_path):
    with open(sql_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    table_name_match = re.search(r'CREATE TABLE `(\w+)`', content)
    table_name = table_name_match.group(1) if table_name_match else 'Unknown'

    columns_match = re.search(r'CREATE TABLE `\w+` \((.*?)\) ENGINE', content, re.DOTALL)
    columns_str = columns_match.group(1) if columns_match else ''

    columns_dict = {}
    lines = columns_str.split('\n')
    for line in lines:
        line = line.strip()
        if line and not line.startswith('PRIMARY') and not line.startswith('UNIQUE') and not line.startswith('KEY'):
            col_match = re.match(r'`(\w+)`', line)
            if col_match:
                col_name = col_match.group(1)
                comment_match = re.search(r"COMMENT '(.*?)'", line)
                comment = comment_match.group(1) if comment_match else ''
                columns_dict[col_name] = comment

    insert_pattern = r"INSERT INTO `\w+` \((.*?)\) VALUES \((.*?)\);"
    inserts = re.findall(insert_pattern, content)

    data = []
    for cols_str, vals_str in inserts:
        cols = [c.strip().strip('`') for c in cols_str.split(',')]
        vals = []
        current_val = ''
        in_quotes = False
        for char in vals_str:
            if char == "'" and (not current_val.endswith('\\') or current_val.endswith('\\\\')):
                in_quotes = not in_quotes
            elif char == ',' and not in_quotes:
                vals.append(current_val.strip())
                current_val = ''
                continue
            current_val += char
        if current_val:
            vals.append(current_val.strip())
        
        row = {}
        for i, col in enumerate(cols):
            if i < len(vals):
                val = vals[i]
                if val == 'NULL':
                    row[col] = None
                elif val.startswith("'") and val.endswith("'"):
                    row[col] = val[1:-1]
                else:
                    try:
                        row[col] = int(val)
                    except ValueError:
                        try:
                            row[col] = float(val)
                        except ValueError:
                            row[col] = val
        data.append(row)

    df = pd.DataFrame(data)

    def format_timestamp(ts_str):
        try:
            ts = int(ts_str)
            if ts == 0:
                return ''
            if ts > 10**12:
                dt = datetime.fromtimestamp(ts / 1000)
            else:
                dt = datetime.fromtimestamp(ts)
            return dt.strftime('%Y-%m-%d %H:%M:%S')
        except:
            return ''

    time_fields = [col for col in df.columns if '_time' in col.lower() or col.lower() == 'time']

    for time_field in time_fields:
        df[time_field] = df[time_field].astype(str)
        format_field_name = f'{time_field}_format'
        df[format_field_name] = df[time_field].apply(format_timestamp)
        columns_dict[format_field_name] = f'{columns_dict.get(time_field, time_field)}（格式化）'

    columns = [{'name': col, 'comment': columns_dict.get(col, '')} for col in df.columns]

    return df, columns, table_name, time_fields

def create_excel(df, columns, excel_file):
    with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='数据记录', index=False)
        
        columns_df = pd.DataFrame(columns)
        columns_df.to_excel(writer, sheet_name='字段说明', index=False)

        workbook = writer.book
        sheet1 = workbook['数据记录']
        sheet2 = workbook['字段说明']

        sheet1.insert_rows(2)

        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        comment_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        comment_font = Font(size=10)

        for idx, col in enumerate(columns, start=1):
            cell1 = sheet1.cell(row=1, column=idx)
            cell1.fill = header_fill
            cell1.font = header_font
            cell1.alignment = Alignment(horizontal='center', vertical='center')

            cell2 = sheet1.cell(row=2, column=idx)
            cell2.value = col['comment']
            cell2.fill = comment_fill
            cell2.font = comment_font
            cell2.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        sheet1.row_dimensions[1].height = 25
        sheet1.row_dimensions[2].height = 30

        for sheet in [sheet1, sheet2]:
            for column in sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                sheet.column_dimensions[column_letter].width = adjusted_width

def clean_path(path):
    return path.strip().strip('"').strip("'")

def main():
    print('=' * 50)
    print('SQL转Excel批量转换工具')
    print('=' * 50)
    print()
    
    sql_input = input('请输入SQL文件路径（多个文件用逗号分隔）: ').strip()
    if not sql_input:
        print('未输入任何文件路径')
        return
    
    sql_files = [clean_path(f) for f in sql_input.split(',')]
    
    output_dir = input('请输入输出目录（直接回车默认 /Users/zyb/Downloads）: ').strip()
    if not output_dir:
        output_dir = '/Users/zyb/Downloads'
    else:
        output_dir = clean_path(output_dir)
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    print()
    print('开始转换...')
    print('-' * 50)
    
    success_count = 0
    error_count = 0
    error_files = []

    for sql_file in sql_files:
        if not os.path.exists(sql_file):
            error_count += 1
            error_files.append(f'{sql_file}: 文件不存在')
            print(f'✗ 文件不存在: {sql_file}')
            continue
            
        try:
            file_name = os.path.basename(sql_file)
            table_name = file_name.replace('.sql', '')
            excel_file = os.path.join(output_dir, f'{table_name}.xlsx')
            
            df, columns, _, time_fields = parse_sql_file(sql_file)
            create_excel(df, columns, excel_file)
            
            success_count += 1
            time_info = f' [时间戳字段: {time_fields}]' if time_fields else ''
            print(f'✓ 转换成功: {file_name} -> {table_name}.xlsx{time_info}')
        except Exception as e:
            error_count += 1
            error_files.append(f'{os.path.basename(sql_file)}: {str(e)}')
            print(f'✗ 转换失败: {os.path.basename(sql_file)} - {str(e)}')

    print('-' * 50)
    print(f'转换完成！')
    print(f'成功: {success_count} 个文件')
    if error_count > 0:
        print(f'失败: {error_count} 个文件')
        for err in error_files:
            print(f'  - {err}')
    print()
    print(f'输出目录: {output_dir}')

if __name__ == '__main__':
    main()
