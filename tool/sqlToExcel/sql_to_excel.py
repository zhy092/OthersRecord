import re
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill
import os
import sys

def clean_path(path):
    return path.strip().strip('"').strip("'")

if len(sys.argv) >= 2:
    sql_file = clean_path(sys.argv[1])
else:
    sql_file = input('请输入SQL文件路径: ').strip()
    
if len(sys.argv) >= 3:
    output_dir = clean_path(sys.argv[2])
else:
    output_dir = input('请输入输出目录（直接回车默认 /Users/zyb/Downloads）: ').strip()
    if not output_dir:
        output_dir = '/Users/zyb/Downloads'

if not os.path.exists(output_dir):
    os.makedirs(output_dir)

with open(sql_file, 'r', encoding='utf-8') as f:
    content = f.read()

table_name_match = re.search(r'CREATE TABLE `(\w+)`', content)
table_name = table_name_match.group(1) if table_name_match else 'Unknown'

columns_match = re.search(r'CREATE TABLE `\w+` \((.*?)\) ENGINE', content, re.DOTALL)
columns_str = columns_match.group(1) if columns_match else ''

columns_dict = {}
lines = columns_str.split('\n')
for line in lines:
    line = line.strip()
    if line and not line.startswith('PRIMARY') and not line.startswith('UNIQUE') and not line.startswith('KEY'):
        col_match = re.match(r'`(\w+)`', line)
        if col_match:
            col_name = col_match.group(1)
            comment_match = re.search(r"COMMENT '(.*?)'", line)
            comment = comment_match.group(1) if comment_match else ''
            columns_dict[col_name] = comment

insert_pattern = r"INSERT INTO `\w+` \((.*?)\) VALUES \((.*?)\);"
inserts = re.findall(insert_pattern, content)

data = []
for cols_str, vals_str in inserts:
    cols = [c.strip().strip('`') for c in cols_str.split(',')]
    vals = []
    current_val = ''
    in_quotes = False
    for char in vals_str:
        if char == "'" and (not current_val.endswith('\\') or current_val.endswith('\\\\')):
            in_quotes = not in_quotes
        elif char == ',' and not in_quotes:
            vals.append(current_val.strip())
            current_val = ''
            continue
        current_val += char
    if current_val:
        vals.append(current_val.strip())
    
    row = {}
    for i, col in enumerate(cols):
        if i < len(vals):
            val = vals[i]
            if val == 'NULL':
                row[col] = None
            elif val.startswith("'") and val.endswith("'"):
                row[col] = val[1:-1]
            else:
                try:
                    row[col] = int(val)
                except ValueError:
                    try:
                        row[col] = float(val)
                    except ValueError:
                        row[col] = val
    data.append(row)

df = pd.DataFrame(data)

def format_timestamp(ts_str):
    try:
        ts = int(ts_str)
        if ts == 0:
            return ''
        if ts > 10**12:
            dt = datetime.fromtimestamp(ts / 1000)
        else:
            dt = datetime.fromtimestamp(ts)
        return dt.strftime('%Y-%m-%d %H:%M:%S')
    except:
        return ''

time_fields = [col for col in df.columns if '_time' in col.lower() or col.lower() == 'time']

for time_field in time_fields:
    df[time_field] = df[time_field].astype(str)
    format_field_name = f'{time_field}_format'
    df[format_field_name] = df[time_field].apply(format_timestamp)
    columns_dict[format_field_name] = f'{columns_dict.get(time_field, time_field)}（格式化）'

columns = [{'name': col, 'comment': columns_dict.get(col, '')} for col in df.columns]

sql_file_name = os.path.basename(sql_file)
excel_name = sql_file_name.replace('.sql', '.xlsx')
excel_file = os.path.join(output_dir, excel_name)

with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='数据记录', index=False)
    
    columns_df = pd.DataFrame(columns)
    columns_df.to_excel(writer, sheet_name='字段说明', index=False)

    workbook = writer.book
    sheet1 = workbook['数据记录']
    sheet2 = workbook['字段说明']

    sheet1.insert_rows(2)

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    comment_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    comment_font = Font(size=10)

    for idx, col in enumerate(columns, start=1):
        cell1 = sheet1.cell(row=1, column=idx)
        cell1.fill = header_fill
        cell1.font = header_font
        cell1.alignment = Alignment(horizontal='center', vertical='center')

        cell2 = sheet1.cell(row=2, column=idx)
        cell2.value = col['comment']
        cell2.fill = comment_fill
        cell2.font = comment_font
        cell2.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    sheet1.row_dimensions[1].height = 25
    sheet1.row_dimensions[2].height = 30

    for sheet in [sheet1, sheet2]:
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width

print(f'Excel文件已生成: {excel_file}')
print(f'共 {len(data)} 条记录')
print(f'共 {len(columns)} 个字段')
if time_fields:
    print(f'自动识别时间戳字段: {time_fields}')
